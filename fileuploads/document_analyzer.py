import os
import re
import json
import joblib
import pandas as pd
import docx
import PyPDF2
import xlrd
from openpyxl import load_workbook

# OCR y extracción avanzada de PDFs
import cv2
import numpy as np
import pytesseract
import pdfplumber
from pdf2image import convert_from_path
import easyocr

# Hugging Face para recomendaciones
from transformers import pipeline

# =============================
# EXTRACCIÓN DE TEXTO MEJORADA
# =============================

class PDFTextExtractor:
    def __init__(self, ocr_engine: str = "pytesseract", lang: str = "spa"):
        self.ocr_engine = ocr_engine
        self.lang = lang
        if ocr_engine == "easyocr":
            lang_code = "es" if lang == "spa" else lang
            self.reader = easyocr.Reader([lang_code])

    def extract_text_from_pdf(self, pdf_path: str) -> str:
        text = self._extract_native_text(pdf_path)
        if not text.strip():
            text = self._extract_scanned_text(pdf_path)
        return text

    def _extract_native_text(self, pdf_path: str) -> str:
        text = ""
        try:
            with pdfplumber.open(pdf_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text(layout=True, x_tolerance=2)
                    if page_text:
                        text += f"\n--- Página {page.page_number} ---\n{page_text}\n"
        except Exception:
            pass
        return text

    def _extract_scanned_text(self, pdf_path: str) -> str:
        images = convert_from_path(pdf_path, dpi=300)
        final_text = ""
        for i, img in enumerate(images):
            img_cv = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
            gray = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)
            _, thresh = cv2.threshold(gray, 150, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
            if self.ocr_engine == "pytesseract":
                page_text = pytesseract.image_to_string(thresh, lang=self.lang, config='--psm 6')
            else:
                results = self.reader.readtext(thresh, paragraph=True)
                page_text = " ".join([res[1] for res in results])
            final_text += f"\n--- Página {i+1} ---\n{page_text}\n"
        return final_text

def extract_text_from_file_advanced(filepath, ocr_engine="pytesseract", lang="spa"):
    """
    Versión mejorada de extracción de texto que incluye OCR
    """
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".txt":
            with open(filepath, "r", encoding="utf-8") as f:
                return f.read()
        elif ext == ".csv":
            return pd.read_csv(filepath).to_string()
        elif ext == ".docx":
            doc = docx.Document(filepath)
            return "\n".join([para.text for para in doc.paragraphs])
        elif ext == ".pdf":
            extractor = PDFTextExtractor(ocr_engine=ocr_engine, lang=lang)
            return extractor.extract_text_from_pdf(filepath)
        elif ext == ".xlsx":
            wb = load_workbook(filepath, read_only=True)
            return "\n".join(str(cell) for sheet in wb for row in sheet.values for cell in row)
        elif ext == ".xls":
            wb = xlrd.open_workbook(filepath)
            return "\n".join(str(cell) for sheet in wb.sheets() for row in sheet.get_rows() for cell in row)
        elif ext in (".json", ".geojson"):
            with open(filepath, "r", encoding="utf-8") as f:
                return json.dumps(json.load(f), ensure_ascii=False)
        else:
            return ""
    except Exception as e:
        print(f"Error en {filepath}: {str(e)[:100]}...")
        return ""

# =============================
# PREPROCESAMIENTO Y CLASIFICACIÓN
# =============================

def preprocess_text(text):
    text = str(text).lower()
    text = re.sub(r'\d+', '', text)
    text = re.sub(r'[^\w\s]', '', text)
    return text.strip()

def classify_document(text):
    
    model = joblib.load("models/document_classifier.joblib")
    vectorizer = joblib.load("models/tfidf_vectorizer.joblib")  
    label_encoder = joblib.load("models/label_encoder.joblib")

    cleaned = preprocess_text(text)
    vectorized = vectorizer.transform([cleaned])
    pred = model.predict(vectorized)[0]
    label = label_encoder.inverse_transform([pred])[0]
    return label

# =============================
# RECOMENDACIÓN DE SALIDA
# =============================

classifier = pipeline("zero-shot-classification", model="facebook/bart-large-mnli")

palabras_clave = {
    "gráfica": ["gráfica", "gráfico", "tendencia", "variación", "cambio porcentual", "tiempo", "comparar", "evolución"],
    "mapa": ["mapa", "ubicación", "localización", "zona", "espacial", "geográfica", "donde", "lugar"],
    "tabla": ["tabla", "datos", "valores", "estadística", "lista", "cantidad", "número"],
    "CSV": ["csv", "datos brutos", "exportar", "procesar", "base de datos"],
    "texto": ["qué", "cómo", "por qué", "metodología", "explicación", "descripción", "resultado", "conclusión"],
    "dashboard": ["resumen", "vista general", "interactivo", "panel"],
    "imagen": ["figura", "foto", "fotografía", "imagen", "ilustración", "visual"]
}

def recomendar_salida(prompt, contexto=""):
    entrada = f"{prompt}\n\nContexto:\n{contexto}" if contexto else prompt
    entrada_proc = preprocess_text(entrada)

    for tipo, palabras in palabras_clave.items():
        for palabra in palabras:
            if palabra in entrada_proc:
                return tipo

    etiquetas = list(palabras_clave.keys())
    resultado = classifier(entrada, candidate_labels=etiquetas)
    return resultado["labels"][0]

# =============================
# FLUJO COMPLETO DE ANÁLISIS
# =============================

def analizar_documento_completo(filepath, prompt=""):
    """
    Función principal para analizar documentos de forma independiente
    """
    texto = extract_text_from_file_advanced(filepath)
    if not texto.strip():
        return {"error": "No se pudo extraer texto"}

    tipo = classify_document(texto)
    recomendacion = recomendar_salida(prompt, texto)

    return {
        "documento": os.path.basename(filepath),
        "tipo": tipo,
        "recomendacion": recomendacion,
        "texto_extraido": texto[:500] + "..." if len(texto) > 500 else texto  # Primeros 500 caracteres
    }

# =============================
# FUNCIÓN PARA USAR CON ARCHIVOS DE DJANGO
# =============================

def analizar_archivo_django(file_object, filename, prompt=""):
    """
    Versión para trabajar directamente con archivos subidos en Django
    """
    temp_path = f"/tmp/{filename}"
    with open(temp_path, 'wb+') as destination:
        for chunk in file_object.chunks():
            destination.write(chunk)
    
    resultado = analizar_documento_completo(temp_path, prompt)
    
    os.remove(temp_path)
    
    return resultado